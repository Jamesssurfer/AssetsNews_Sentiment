"""Rebuild the 'Daily' summary sheet from the 'Log' sheet: one row per
(date, bucket), with a live Excel formula for day-over-day % change. Run
once a day via .github/workflows/daily_rollup.yml.

The Daily sheet is fully rebuilt each run from Log -- it's a derived view,
not a place to hand-edit.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

LOG_PATH = Path(__file__).parent / "data" / "asset_sentiment_log.xlsx"
LOG_SHEET = "Log"
DAILY_SHEET = "Daily"

DAILY_HEADERS = [
    "date", "bucket", "bullish_articles", "bearish_articles",
    "composite_avg", "composite_pct_change_vs_prev_day",
    "ohlc_close_avg",
]


def build_daily_frame() -> pd.DataFrame:
    df = pd.read_excel(LOG_PATH, sheet_name=LOG_SHEET)
    grouped = (
        df.groupby(["date", "bucket"], as_index=False)
        .agg(
            bullish_articles=("bullish_articles", "sum"),
            bearish_articles=("bearish_articles", "sum"),
            composite_avg=("composite_score", "mean"),
            ohlc_close_avg=("ohlc_close", "mean"),
        )
    )
    grouped["composite_avg"] = grouped["composite_avg"].round(3)
    grouped["ohlc_close_avg"] = grouped["ohlc_close_avg"].round(4)
    grouped = grouped.sort_values(["bucket", "date"]).reset_index(drop=True)
    return grouped


def write_daily_sheet(grouped: pd.DataFrame):
    wb = load_workbook(LOG_PATH)
    if DAILY_SHEET in wb.sheetnames:
        del wb[DAILY_SHEET]
    ws = wb.create_sheet(DAILY_SHEET)
    ws.append(DAILY_HEADERS)

    for i, row in enumerate(grouped.itertuples(index=False), start=2):  # row 1 = header
        ws.append([
            row.date, row.bucket, row.bullish_articles, row.bearish_articles,
            row.composite_avg, None, row.ohlc_close_avg,
        ])
        formula = f'=IF(B{i-1}=B{i}, IFERROR((E{i}-E{i-1})/ABS(E{i-1}), "n/a"), "")'
        ws[f"F{i}"] = formula

    wb.save(LOG_PATH)


def run_rollup():
    if not LOG_PATH.exists():
        print(f"No log file yet at {LOG_PATH} -- run scanner.py at least once first.")
        return
    grouped = build_daily_frame()
    write_daily_sheet(grouped)
    print(f"Daily sheet rebuilt: {len(grouped)} (date, bucket) rows.")


if __name__ == "__main__":
    run_rollup()
