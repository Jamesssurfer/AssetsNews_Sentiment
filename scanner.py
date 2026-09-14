"""One scan cycle: for each asset bucket, fetch bullish + bearish news, score
it, fetch the trailing 4-hour OHLC price window, and append one row per
bucket to the running Excel log.

Run manually with: python scanner.py
Run on a schedule via .github/workflows/scan.yml (every 4 hours).
"""
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config import BUCKETS, MAX_ARTICLES_PER_QUERY
from news_fetcher import fetch_and_score_group
from sentiment_engine import score_text
import price_fetcher

LOG_PATH = Path(__file__).parent / "data" / "asset_sentiment_log.xlsx"
LOG_SHEET = "Log"

HEADERS = [
    "timestamp_utc", "date", "bucket",
    "ticker_used", "ohlc_open", "ohlc_high", "ohlc_low", "ohlc_close",
    "ohlc_volume", "ohlc_bars_used",
    "bullish_articles", "bullish_finbert_avg", "bullish_vader_avg",
    "bearish_articles", "bearish_finbert_avg", "bearish_vader_avg",
    "composite_score", "sentiment_model_used",
    "bullish_top_headline_1", "bullish_top_headline_1_url",
    "bullish_top_headline_2", "bullish_top_headline_2_url",
    "bullish_top_headline_3", "bullish_top_headline_3_url",
    "bearish_top_headline_1", "bearish_top_headline_1_url",
    "bearish_top_headline_2", "bearish_top_headline_2_url",
    "bearish_top_headline_3", "bearish_top_headline_3_url",
]


def _avg(values: list[float]) -> float:
    return sum(values) / len(values) if values else 0.0


def _composite_score(bullish_count, bullish_score_avg, bearish_count, bearish_score_avg) -> float:
    """
    Positive = net bullish pressure. Negative = net bearish pressure. Zero = quiet on both sides.

    Mirrors the geopolitical repo's risk/de-escalation formula shape, with the
    roles reversed to get the sign convention right for this repo:
      bullish_weight = bullish_count * (0.5 + bullish_score_avg / 2)
                        (more articles + more positive tone -> higher)
      bearish_weight = bearish_count * (0.5 - bearish_score_avg / 2)
                        (more articles + more negative tone -> higher)
      composite      = bullish_weight - bearish_weight

    v1 heuristic, not a validated model -- same caveat as the other repo.
    """
    bullish_weight = bullish_count * (0.5 + bullish_score_avg / 2)
    bearish_weight = bearish_count * (0.5 - bearish_score_avg / 2)
    return round(bullish_weight - bearish_weight, 3)


def _top_headlines(articles: list[dict], score_key: str, n: int = 3) -> list[dict]:
    ranked = sorted(articles, key=lambda a: abs(a[score_key]), reverse=True)[:n]
    padded = ranked + [{"title": "", "link": ""}] * (n - len(ranked))
    return padded


def _migrate_headers(ws):
    """Same append-only migration discipline as the geopolitical repo: never
    insert a new column mid-schema (that desyncs every row already on disk),
    only ever extend the header row with columns HEADERS adds at the end."""
    existing = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    if existing == HEADERS:
        return
    if existing and HEADERS[: len(existing)] == existing:
        for i, header in enumerate(HEADERS[len(existing):], start=len(existing) + 1):
            ws.cell(row=1, column=i, value=header)
    else:
        raise RuntimeError(
            "data/asset_sentiment_log.xlsx header row doesn't match the expected schema "
            "-- this needs a manual look before more rows get appended.\n"
            f"  Found:    {existing}\n"
            f"  Expected: {HEADERS}"
        )


def _ensure_workbook():
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    if LOG_PATH.exists():
        wb = load_workbook(LOG_PATH)
        if LOG_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(LOG_SHEET)
            ws.append(HEADERS)
        else:
            _migrate_headers(wb[LOG_SHEET])
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = LOG_SHEET
        ws.append(HEADERS)
    return wb


def run_scan():
    now = datetime.now(timezone.utc)
    timestamp = now.isoformat(timespec="seconds")
    date_str = now.strftime("%Y-%m-%d")

    wb = _ensure_workbook()
    ws = wb[LOG_SHEET]

    price_session = price_fetcher.make_session()

    for bucket_key, bucket in BUCKETS.items():
        model = bucket.get("sentiment_model", "finbert")
        print(f"Scanning bucket: {bucket['label']} (scoring model: {model})")

        bullish_articles = fetch_and_score_group(bucket["bullish_queries"], MAX_ARTICLES_PER_QUERY, score_text)
        bearish_articles = fetch_and_score_group(bucket["bearish_queries"], MAX_ARTICLES_PER_QUERY, score_text)

        bullish_finbert_avg = _avg([a["finbert"] for a in bullish_articles])
        bullish_vader_avg = _avg([a["vader"] for a in bullish_articles])
        bearish_finbert_avg = _avg([a["finbert"] for a in bearish_articles])
        bearish_vader_avg = _avg([a["vader"] for a in bearish_articles])

        bullish_active_avg = bullish_vader_avg if model == "vader" else bullish_finbert_avg
        bearish_active_avg = bearish_vader_avg if model == "vader" else bearish_finbert_avg

        composite = _composite_score(
            len(bullish_articles), bullish_active_avg,
            len(bearish_articles), bearish_active_avg,
        )

        bullish_top3 = _top_headlines(bullish_articles, model, 3)
        bearish_top3 = _top_headlines(bearish_articles, model, 3)

        price = price_fetcher.fetch_bucket_price(bucket, price_session)

        row = [
            timestamp, date_str, bucket["label"],
            price["ticker_used"], price["ohlc_open"], price["ohlc_high"],
            price["ohlc_low"], price["ohlc_close"], price["ohlc_volume"], price["ohlc_bars_used"],
            len(bullish_articles), round(bullish_finbert_avg, 3), round(bullish_vader_avg, 3),
            len(bearish_articles), round(bearish_finbert_avg, 3), round(bearish_vader_avg, 3),
            composite, model,
        ]
        for h in bullish_top3:
            row.extend([h.get("title", ""), h.get("link", "")])
        for h in bearish_top3:
            row.extend([h.get("title", ""), h.get("link", "")])

        ws.append(row)
        print(
            f"  bullish={len(bullish_articles)} articles (active tone {bullish_active_avg:.2f}), "
            f"bearish={len(bearish_articles)} articles (active tone {bearish_active_avg:.2f}), "
            f"composite={composite}, price={price['ticker_used']} close={price['ohlc_close']}"
        )

    wb.save(LOG_PATH)
    print(f"Saved log -> {LOG_PATH}")


if __name__ == "__main__":
    run_scan()
